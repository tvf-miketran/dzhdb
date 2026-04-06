from __future__ import annotations

import io
import json
import re
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any, Dict, List, Tuple
from uuid import UUID

from openpyxl import Workbook, load_workbook
from sqlalchemy import inspect, text
from sqlalchemy.sql.sqltypes import ARRAY as SQLArray
from sqlalchemy.sql.sqltypes import JSON as SQLJson

from app import db


class ExportService:
    EXCLUDED_TABLES = {"alembic_version"}

    @staticmethod
    def _get_exportable_table_names() -> List[str]:
        inspector = inspect(db.engine)
        all_tables = sorted(inspector.get_table_names())
        return [
            table_name
            for table_name in all_tables
            if ExportService._is_exportable_table(table_name)
        ]

    @staticmethod
    def get_exportable_tables() -> Dict[str, Any]:
        inspector = inspect(db.engine)
        all_tables = sorted(inspector.get_table_names())

        exportable_tables: List[Dict[str, Any]] = []
        excluded_tables: List[str] = []

        for table_name in all_tables:
            if not ExportService._is_exportable_table(table_name):
                excluded_tables.append(table_name)
                continue

            columns = [col["name"] for col in inspector.get_columns(table_name)]
            exportable_tables.append(
                {
                    "table": table_name,
                    "columns": columns,
                }
            )

        return {
            "tables": exportable_tables,
            "table_names": [item["table"] for item in exportable_tables],
            "excluded_tables": excluded_tables,
            "total": len(exportable_tables),
        }

    @staticmethod
    def export_database_to_excel() -> Tuple[io.BytesIO, Dict[str, Any]]:
        inspector = inspect(db.engine)
        table_names = ExportService._get_exportable_table_names()

        if not table_names:
            raise ValueError("No tables to export")

        workbook = Workbook(write_only=True)
        sheet_name_registry: set[str] = set()
        sheet_row_counts: Dict[str, int] = {}

        for table_name in table_names:
            table_columns = [col["name"] for col in inspector.get_columns(table_name)]
            selected_columns = table_columns
            select_columns_sql = ", ".join([f'"{col}"' for col in table_columns])
            sql = f'SELECT {select_columns_sql} FROM "{table_name}"'

            rows = db.session.execute(text(sql)).mappings().all()

            sheet_name = ExportService._safe_sheet_name(table_name, sheet_name_registry)
            worksheet = workbook.create_sheet(title=sheet_name)
            worksheet.append(selected_columns)

            for row in rows:
                worksheet.append(
                    [ExportService._serialize_excel_value(row.get(col)) for col in selected_columns]
                )

            sheet_row_counts[table_name] = len(rows)

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        metadata = {
            "table_count": len(table_names),
            "tables": table_names,
            "row_count_by_table": sheet_row_counts,
            "exported_at": datetime.utcnow().isoformat(),
        }
        return output, metadata

    @staticmethod
    def import_database_from_excel(file_stream: Any) -> Dict[str, Any]:
        workbook = load_workbook(filename=file_stream, read_only=True, data_only=True)
        inspector = inspect(db.engine)

        expected_tables = ExportService._get_exportable_table_names()
        exportable_tables = set(expected_tables)
        workbook_sheets = workbook.sheetnames

        workbook_sheet_set = set(workbook_sheets)
        missing_tables = sorted(exportable_tables - workbook_sheet_set)
        unexpected_tables = sorted(workbook_sheet_set - exportable_tables)

        if missing_tables or unexpected_tables:
            raise ValueError(
                "Invalid import file format. "
                f"missing tables: {missing_tables}, unexpected sheets: {unexpected_tables}"
            )

        table_names = expected_tables

        table_payloads: Dict[str, Dict[str, Any]] = {}

        for table_name in table_names:
            worksheet = workbook[table_name]
            rows_iter = worksheet.iter_rows(values_only=True)
            header_row = next(rows_iter, None)

            if not header_row:
                raise ValueError(f"Sheet '{table_name}' does not contain a header row")

            db_columns = inspector.get_columns(table_name)
            expected_columns = [col["name"] for col in db_columns]
            column_type_map = {col["name"]: col.get("type") for col in db_columns}
            header_len = len(expected_columns)
            header_values = list(header_row[:header_len])
            if len(header_values) < header_len:
                header_values.extend([None] * (header_len - len(header_values)))

            headers = [str(col).strip() if col is not None else "" for col in header_values]
            if headers != expected_columns:
                raise ValueError(
                    f"Invalid header format in sheet '{table_name}'. "
                    f"Expected: {expected_columns}, got: {headers}"
                )

            records: List[Dict[str, Any]] = []

            for row in rows_iter:
                if not row:
                    continue

                row_values = list(row[:header_len])
                if len(row_values) < header_len:
                    row_values.extend([None] * (header_len - len(row_values)))

                if all(value is None or (isinstance(value, str) and value.strip() == "") for value in row_values):
                    continue

                record = {
                    expected_columns[idx]: ExportService._deserialize_excel_value(
                        row_values[idx],
                        column_type_map.get(expected_columns[idx]),
                    )
                    for idx in range(header_len)
                }
                records.append(record)

            table_payloads[table_name] = {
                "columns": expected_columns,
                "records": records,
            }

        insert_order = ExportService._resolve_insert_order(table_names, inspector)
        truncate_tables_sql = ", ".join([f'"{table_name}"' for table_name in sorted(exportable_tables)])
        inserted_row_count: Dict[str, int] = {}

        # Keep import atomic while avoiding nested transaction conflicts.
        try:
            db.session.rollback()
            db.session.execute(text(f"TRUNCATE TABLE {truncate_tables_sql} CASCADE"))

            for table_name in insert_order:
                payload = table_payloads.get(table_name) or {}
                columns = payload.get("columns") or []
                records = payload.get("records") or []

                if not columns or not records:
                    inserted_row_count[table_name] = 0
                    continue

                column_sql = ", ".join([f'"{column}"' for column in columns])
                value_sql = ", ".join([f":{column}" for column in columns])
                insert_stmt = text(
                    f'INSERT INTO "{table_name}" ({column_sql}) VALUES ({value_sql})'
                )

                db.session.execute(insert_stmt, records)
                inserted_row_count[table_name] = len(records)

            db.session.commit()
        except Exception:
            db.session.rollback()
            raise

        return {
            "imported_tables": insert_order,
            "imported_table_count": len(insert_order),
            "inserted_row_count_by_table": inserted_row_count,
            "total_inserted_rows": sum(inserted_row_count.values()),
            "cleaned_tables_before_import": sorted(exportable_tables),
            "cleaned_table_count": len(exportable_tables),
            "imported_at": datetime.utcnow().isoformat(),
        }

    @staticmethod
    def _is_exportable_table(table_name: str) -> bool:
        lower_name = table_name.lower()
        if table_name in ExportService.EXCLUDED_TABLES:
            return False
        if "migration" in lower_name:
            return False
        return True

    @staticmethod
    def _safe_sheet_name(original_name: str, used_names: set[str]) -> str:
        safe_name = re.sub(r"[\\/*?:\[\]]", "_", original_name)[:31]
        if not safe_name:
            safe_name = "sheet"

        candidate = safe_name
        suffix = 1
        while candidate.lower() in used_names:
            tail = f"_{suffix}"
            allowed = 31 - len(tail)
            candidate = f"{safe_name[:allowed]}{tail}"
            suffix += 1

        used_names.add(candidate.lower())
        return candidate

    @staticmethod
    def _serialize_excel_value(value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, (datetime, date, time)):
            return value.isoformat()
        if isinstance(value, UUID):
            return str(value)
        if isinstance(value, Decimal):
            return float(value)
        if isinstance(value, (dict, list)):
            return json.dumps(value, ensure_ascii=True)
        return value

    @staticmethod
    def _deserialize_excel_value(value: Any, column_type: Any = None) -> Any:
        if isinstance(value, str):
            stripped_value = value.strip()
            if stripped_value == "":
                return None

            if isinstance(column_type, SQLArray):
                if stripped_value.startswith("[") and stripped_value.endswith("]"):
                    try:
                        parsed = json.loads(stripped_value)
                        if isinstance(parsed, list):
                            return parsed
                    except json.JSONDecodeError:
                        pass

            if isinstance(column_type, SQLJson):
                if (
                    (stripped_value.startswith("{") and stripped_value.endswith("}"))
                    or (stripped_value.startswith("[") and stripped_value.endswith("]"))
                ):
                    try:
                        return json.loads(stripped_value)
                    except json.JSONDecodeError:
                        pass

            return stripped_value
        return value

    @staticmethod
    def _resolve_insert_order(table_names: List[str], inspector: Any) -> List[str]:
        table_set = set(table_names)
        dependencies: Dict[str, set[str]] = {table: set() for table in table_names}

        for table in table_names:
            for fk in inspector.get_foreign_keys(table):
                referred_table = fk.get("referred_table")
                if referred_table in table_set and referred_table != table:
                    dependencies[table].add(referred_table)

        indegree: Dict[str, int] = {table: len(dependencies[table]) for table in table_names}
        queue = sorted([table for table, degree in indegree.items() if degree == 0])
        order: List[str] = []

        while queue:
            current = queue.pop(0)
            order.append(current)

            for table in table_names:
                if current in dependencies[table]:
                    dependencies[table].remove(current)
                    indegree[table] -= 1
                    if indegree[table] == 0:
                        queue.append(table)
                        queue.sort()

        if len(order) < len(table_names):
            remaining = [table for table in table_names if table not in order]
            order.extend(sorted(remaining))

        return order
